# -*- coding: utf-8 -*-
"""Build small sample .pptx decks for tests and demos.

Needs python-pptx (development only).  The app itself parses PPTX with the
standard library, so this script is never shipped.

    python tools/make_sample_pptx.py
"""

from __future__ import annotations

import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:  # pragma: no cover
    print("需要 python-pptx：pip install python-pptx")
    sys.exit(2)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "samples", "课件示例")

DECKS = [
    (
        "第1讲 极限与连续.pptx",
        [
            ("1.1 数列的极限", ["数列：按正整数编号的一列数", "研究 n 越来越大时 a_n 的趋势"], ""),
            ("定义 1.1（ε-N 定义）", ["对任意 ε>0，存在 N，当 n>N 时 |a_n - A| < ε", "记作 lim a_n = A"], "强调 N 依赖于 ε，但不要求最小"),
            ("例 1.1", ["用定义证明 lim (n+1)/n = 1", "关键：|(n+1)/n - 1| = 1/n"], ""),
            ("1.2 收敛数列的性质", ["定理 1.1 唯一性", "定理 1.2 有界性", "定理 1.3 保号性"], ""),
            ("小结", ["定义是证明的工具", "性质是计算的依据"], ""),
        ],
    ),
    (
        "第2讲 导数与微分.pptx",
        [
            ("2.1 导数的定义", ["差商的极限", "f'(x0) = lim [f(x0+Δx)-f(x0)]/Δx"], "导数刻画瞬时变化率"),
            ("例 2.1", ["求 f(x)=x^2 在 x=1 处的导数", "先写差商，再约分，再取极限"], ""),
            ("2.2 求导法则", ["(uv)' = u'v + uv'", "(u/v)' = (u'v - uv')/v^2"], ""),
            ("2.3 高阶导数", ["f''(x) 是 f'(x) 的导数", "莱布尼茨公式"], ""),
        ],
    ),
    (
        "第10讲 微分方程初步.pptx",
        [
            ("10.1 基本概念", ["含有未知函数及其导数的方程叫微分方程", "阶、通解、特解"], ""),
            ("10.2 可分离变量方程", ["dy/dx = f(x)g(y)", "分离后两边积分"], ""),
        ],
    ),
]


def build_one(path: str, title: str, slides):
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    first = prs.slides.add_slide(blank)
    box = first.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(8.4), Inches(1.4))
    tf = box.text_frame
    tf.text = title.replace(".pptx", "")
    tf.paragraphs[0].runs[0].font.size = Pt(36)
    tf.paragraphs[0].runs[0].font.bold = True

    for heading, bullets, notes in slides:
        slide = prs.slides.add_slide(blank)
        head = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.6), Inches(0.9))
        hp = head.text_frame.paragraphs[0]
        hp.text = heading
        hp.runs[0].font.size = Pt(28)
        hp.runs[0].font.bold = True

        body = slide.shapes.add_textbox(Inches(0.9), Inches(1.6), Inches(8.2), Inches(4.6))
        tf2 = body.text_frame
        tf2.word_wrap = True
        for i, line in enumerate(bullets):
            p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
            p.text = "• " + line
            p.runs[0].font.size = Pt(20)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(path)
    return os.path.getsize(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    for name, slides in DECKS:
        path = os.path.join(OUT, name)
        size = build_one(path, name, slides)
        print("wrote %s (%d bytes, %d slides)" % (name, size, len(slides) + 1))
    print("out dir:", OUT)


if __name__ == "__main__":
    main()
